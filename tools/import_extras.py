# -*- coding: utf-8 -*-
"""
Extras -> AddOn PaymentPlan + MainAddOnLink (zorunlular için).
Strateji:
- Category="Accommodation Placement" -> ATLA (accommodation import'ta zaten kullanıldı)
- Category="Registration"           -> ilgili (school,campus) Main'lerinin RegistrationFee'sine yaz
- Category="Country Note" / RequiredStatus="Note" -> ATLA (bilgi, ücret değil)
- Diğerleri (Transfer/Material/Insurance/...) -> AddOn PaymentPlan
- Mandatory / Mandatory for Student Visa -> IsOrHasMandatory=1 + o (school,campus) Main'lerine link
"""
import openpyxl, pymysql

PATH = r"C:\Users\gorkem.turhan\Downloads\PİVOT SON DOĞRU OKUL.xlsx"
conn = pymysql.connect(host="127.0.0.1", port=3306, user="root", password="Pivot.2026!",
                       database="pivot", charset="utf8mb4", autocommit=False)
cur = conn.cursor()

wb = openpyxl.load_workbook(PATH, data_only=True)
ws = wb["Extras"]
h = [c.value for c in ws[1]]
xi = {x: i for i, x in enumerate(h) if x is not None}

# School lookup (Country, City, School) -> Id (Course_Prices ve Accommodation import'tan)
cur.execute("""SELECT s.Id, s.Name, ci.Name, co.Name
               FROM Schools s
               JOIN Cities ci ON ci.Id=s.CityId
               JOIN Countries co ON co.Id=ci.CountryId""")
school_by_key = {(co, ci, sn): sid for sid, sn, ci, co in cur.fetchall()}

def get_school_id(country, city, school):
    return school_by_key.get((country, city, school))

# Bu (school,campus)'ün TÜM Main PaymentPlan id'lerini getiren cache
mains_cache = {}
def get_main_ids(sid):
    if sid in mains_cache: return mains_cache[sid]
    cur.execute("SELECT Id FROM PaymentPlans WHERE SchoolId=%s AND PackageType=1", (sid,))
    ids = [row[0] for row in cur.fetchall()]
    mains_cache[sid] = ids
    return ids

reg_updates = 0
addon_inserts = 0
link_inserts = 0
skipped = 0

addon_sql = """INSERT INTO PaymentPlans
(SchoolId, ProgramId, Name, MinWeek, MaxWeek, PriceType,
 WeeklyListFee, TotalListFee, PackageType, IsAdditional,
 Category, IsOrHasMandatory, IsActive)
VALUES (%s, NULL, %s, 1, 999, %s, %s, %s, 2, 1, %s, %s, 1)"""

link_sql = "INSERT INTO MainAddOnLinks (MainPlanId, AddOnPlanId) VALUES (%s, %s)"

for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None or str(r[0]).strip() == "": continue
    school = str(r[xi["School"]]).strip()
    country = str(r[xi["Country"]]).strip()
    campus = str(r[xi["Campus"]]).strip()
    fee_name = str(r[xi["FeeName"]] or "").strip()
    category = str(r[xi["Category"]] or "").strip()
    fee_type = str(r[xi["FeeType"]] or "").strip()
    amount = r[xi["Amount"]]
    req = str(r[xi["RequiredStatus"]] or "").strip()

    if category == "Accommodation Placement":
        skipped += 1; continue
    if category == "Country Note" or req == "Note":
        skipped += 1; continue

    sid = get_school_id(country, campus, school)
    if sid is None:
        skipped += 1; continue

    if category == "Registration":
        # UK için 0 olabilir; yine de yazalım (Cart 0 ise göstermez).
        amt = float(amount) if amount is not None else 0.0
        mains = get_main_ids(sid)
        if mains:
            fmt = ",".join(["%s"] * len(mains))
            cur.execute(f"UPDATE PaymentPlans SET RegistrationFee=%s WHERE Id IN ({fmt})",
                        [amt] + mains)
            reg_updates += len(mains)
        continue

    # AddOn olarak ekle
    is_mandatory = 1 if req in ("Mandatory", "Mandatory for Student Visa") else 0
    if fee_type == "Weekly":
        price_type = 1
        weekly = float(amount) if amount is not None else 0.0
        total = None
    else:
        # Once, FixedTotal, GreenwichOSHC, LexisMaterials, Monthly, WeeklyMinMax, WeeklyCapped20, Contact
        # Contact için amount null olabilir; o zaman 0 yaz (UI'da not göstermek lazım).
        price_type = 2
        weekly = None
        total = float(amount) if amount is not None else 0.0

    cur.execute(addon_sql, (sid, fee_name, price_type, weekly, total, category, is_mandatory))
    aid = cur.lastrowid
    addon_inserts += 1

    if is_mandatory:
        for mid in get_main_ids(sid):
            cur.execute(link_sql, (mid, aid))
            link_inserts += 1
        # main'lerin IsOrHasMandatory bayrağını da set et
        mains = get_main_ids(sid)
        if mains:
            fmt = ",".join(["%s"] * len(mains))
            cur.execute(f"UPDATE PaymentPlans SET IsOrHasMandatory=1 WHERE Id IN ({fmt})", mains)

conn.commit()
print(f"Registration Fee güncellenen Main: {reg_updates}")
print(f"AddOn (Transfer/Material/Insurance/...) eklenen: {addon_inserts}")
print(f"MainAddOnLink (zorunlu): {link_inserts}")
print(f"Atlanan satır: {skipped}")
cur.close(); conn.close()
