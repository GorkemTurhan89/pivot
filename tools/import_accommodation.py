# -*- coding: utf-8 -*-
"""
Accommodation_Prices -> AddOn PaymentPlan + AccommodationDetail (1:1).
Ayrıca Extras'tan o (School,Campus) için 'Accommodation Placement Fee' okunup
PaymentPlan.RegistrationFee'ye yazılır — yani konaklama seçilince otomatik eklenir.
Not loaded satırları atlanır.
"""
import openpyxl, pymysql, datetime
from collections import defaultdict

PATH = r"C:\Users\gorkem.turhan\Downloads\PİVOT SON DOĞRU OKUL.xlsx"

conn = pymysql.connect(host="127.0.0.1", port=3306, user="root", password="Pivot.2026!",
                       database="pivot", charset="utf8mb4", autocommit=False)
cur = conn.cursor()

wb = openpyxl.load_workbook(PATH, data_only=True)

# 1) Önce Extras'tan (school, campus) -> placement_fee dict.
ws_ex = wb["Extras"]
hx = [c.value for c in ws_ex[1]]
xi = {h: i for i, h in enumerate(hx) if h is not None}
placement = {}
for r in ws_ex.iter_rows(min_row=2, values_only=True):
    if r[0] is None or str(r[0]).strip() == "": continue
    if str(r[xi["Category"]] or "").strip() != "Accommodation Placement": continue
    s = str(r[xi["School"]]).strip()
    cm = str(r[xi["Campus"]]).strip()
    amt = r[xi["Amount"]]
    if amt is None: continue
    try: placement[(s, cm)] = float(amt)
    except: pass
print(f"Accommodation Placement Fee bulunan kampüsler: {len(placement)}")

# 2) Mevcut School kayıtlarını çek (Course_Prices import'undan kalan).
cur.execute("""SELECT s.Id, s.Name, ci.Name, co.Name
               FROM Schools s
               JOIN Cities ci ON ci.Id=s.CityId
               JOIN Countries co ON co.Id=ci.CountryId""")
school_by_key = {(co, ci, sn): sid for sid, sn, ci, co in cur.fetchall()}

# Country/City/School get-or-create (Accommodation'da yeni okul/şehir çıkabilir).
country_id = {}
def get_country(name):
    if name in country_id: return country_id[name]
    cur.execute("SELECT Id FROM Countries WHERE Name=%s", (name,))
    row = cur.fetchone()
    if row: country_id[name] = row[0]; return row[0]
    # yeni ülke - currency bilinmiyor, 'USD' default ver (sonra düzeltilebilir)
    cur.execute("INSERT INTO Countries (Name, Currency) VALUES (%s, 'USD')", (name,))
    country_id[name] = cur.lastrowid
    return country_id[name]

city_id = {}
def get_city(country, city):
    key = (country, city)
    if key in city_id: return city_id[key]
    cid = get_country(country)
    cur.execute("SELECT Id FROM Cities WHERE Name=%s AND CountryId=%s", (city, cid))
    row = cur.fetchone()
    if row: city_id[key] = row[0]; return row[0]
    cur.execute("INSERT INTO Cities (Name, CountryId) VALUES (%s, %s)", (city, cid))
    city_id[key] = cur.lastrowid
    return city_id[key]

def get_school(country, city, school):
    key = (country, city, school)
    if key in school_by_key: return school_by_key[key]
    cid = get_city(country, city)
    cur.execute("INSERT INTO Schools (Name, CityId) VALUES (%s, %s)", (school, cid))
    sid = cur.lastrowid
    school_by_key[key] = sid
    return sid

# 3) Accommodation_Prices satırlarını işle.
ws = wb["Accommodation_Prices"]
ha = [c.value for c in ws[1]]
ai = {h: i for i, h in enumerate(ha) if h is not None}

plan_sql = """INSERT INTO PaymentPlans
(SchoolId, ProgramId, Name, MinWeek, MaxWeek, PriceType,
 WeeklyListFee, RegistrationFee, PackageType, IsAdditional,
 Category, IsOrHasMandatory, IsActive)
VALUES (%s, NULL, %s, %s, %s, 1, %s, %s, 2, 1, 'Konaklama', 0, 1)"""

detail_sql = """INSERT INTO AccommodationDetails
(PaymentPlanId, Type, RoomType, Board, ResidenceName,
 School, Country, Campus, Required, DisplayOption)
VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)"""

ins = 0; skipped = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None or str(r[0]).strip() == "": continue
    school = str(r[ai["School"]]).strip()
    country = str(r[ai["Country"]]).strip()
    campus = str(r[ai["Campus"]]).strip()
    atype = str(r[ai["AccommodationType"]] or "").strip()
    room = str(r[ai["RoomType"]] or "").strip()
    board = str(r[ai["Board"]] or "").strip()
    resname = r[ai["ResidenceName"]]
    resname = str(resname).strip() if resname else None
    req_status = str(r[ai["RequiredStatus"]] or "").strip()
    weekly = r[ai["WeeklyFee"]]

    if req_status == "Not loaded" or weekly is None:
        skipped += 1; continue

    min_w = int(r[ai["MinWeek"]]) if r[ai["MinWeek"]] is not None else 1
    max_w = int(r[ai["MaxWeek"]]) if r[ai["MaxWeek"]] is not None else 999
    display = r[ai["DisplayOption"]]
    name = str(display).strip() if display else f"{atype} / {room} / {board}"

    sid = get_school(country, campus, school)
    reg_fee = placement.get((school, campus))

    cur.execute(plan_sql, (sid, name, min_w, max_w, float(weekly), reg_fee))
    pid = cur.lastrowid
    cur.execute(detail_sql, (pid, atype, room, board, resname,
                             school, country, campus, req_status, name))
    ins += 1

conn.commit()
print(f"Accommodation PaymentPlan eklendi: {ins} | Skip: {skipped}")
cur.close(); conn.close()
