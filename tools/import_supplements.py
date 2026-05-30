# -*- coding: utf-8 -*-
"""
Supplements -> AddOn PaymentPlan (Category="Supplement") + SupplementDetail (1:1).
Eşleşme runtime'da: AccommodationDetail.(School, Country, Campus, Type) ==
SupplementDetail.(School, Country, Campus, AppliesTo).
Conditional satırlar UI'da kırmızı bar + kilit + popup ile gösterilir.
"""
import openpyxl, pymysql, datetime

PATH = r"C:\Users\gorkem.turhan\Downloads\PİVOT SON DOĞRU OKUL.xlsx"
conn = pymysql.connect(host="127.0.0.1", port=3306, user="root", password="Pivot.2026!",
                       database="pivot", charset="utf8mb4", autocommit=False)
cur = conn.cursor()

wb = openpyxl.load_workbook(PATH, data_only=True)
ws = wb["Supplements"]
h = [c.value for c in ws[1]]
xi = {x: i for i, x in enumerate(h) if x is not None}

cur.execute("""SELECT s.Id, s.Name, ci.Name, co.Name
               FROM Schools s
               JOIN Cities ci ON ci.Id=s.CityId
               JOIN Countries co ON co.Id=ci.CountryId""")
school_by_key = {(co, ci, sn): sid for sid, sn, ci, co in cur.fetchall()}

def to_date(v):
    if v is None or v == "": return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.date() if isinstance(v, datetime.datetime) else v
    try: return datetime.date.fromisoformat(str(v)[:10])
    except: return None

plan_sql = """INSERT INTO PaymentPlans
(SchoolId, ProgramId, Name, MinWeek, MaxWeek, PriceType,
 WeeklyListFee, PackageType, IsAdditional, Category, IsOrHasMandatory, IsActive)
VALUES (%s, NULL, %s, 1, 999, 1, %s, 2, 1, 'Supplement', 0, 1)"""

detail_sql = """INSERT INTO SupplementDetails
(PaymentPlanId, School, Country, Campus, AppliesTo, SupplementName,
 Required, Notes, StartDate, EndDate)
VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)"""

ins = 0; skipped = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None or str(r[0]).strip() == "": continue
    school = str(r[xi["School"]]).strip()
    country = str(r[xi["Country"]]).strip()
    campus = str(r[xi["Campus"]]).strip()
    supp_name = str(r[xi["SupplementName"]] or "").strip()
    applies_to = str(r[xi["AppliesTo"]] or "").strip()
    amount = r[xi["Amount"]]
    req = str(r[xi["RequiredStatus"]] or "").strip()
    notes = r[xi["Notes"]]
    notes = str(notes).strip() if notes else None
    start = to_date(r[xi["StartDate"]])
    end = to_date(r[xi["EndDate"]])

    sid = school_by_key.get((country, campus, school))
    if sid is None or amount is None:
        skipped += 1; continue

    name = f"{supp_name} ({applies_to})"
    cur.execute(plan_sql, (sid, name, float(amount)))
    pid = cur.lastrowid
    cur.execute(detail_sql, (pid, school, country, campus, applies_to,
                             supp_name, req, notes, start, end))
    ins += 1

conn.commit()
print(f"Supplement eklenen: {ins} | Skip: {skipped}")
cur.close(); conn.close()
